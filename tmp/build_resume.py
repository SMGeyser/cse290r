from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\samge\school\homework_f26\cse290r")
OUT = ROOT / "output" / "docx" / "Samuel_Geyser_Northrop_Grumman_Intern_Resume.docx"


def set_font(run, name="Arial", size=10, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_hyperlink(paragraph, text, url, color="1F4E79", underline=True):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(r_fonts)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "18")
    r_pr.append(size_cs)

    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if underline:
        underline_el = OxmlElement("w:u")
        underline_el.set(qn("w:val"), "single")
        r_pr.append(underline_el)

    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph(style="Resume Section")
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_spacing(paragraph, before=4.2, after=1.2, line=1.0)
    run = paragraph.add_run(text.upper())
    set_font(run, size=10.5, bold=True)
    return paragraph


def add_entry_heading(doc, title, dates, organization, location):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=1.5, after=0, line=1.0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.30), WD_TAB_ALIGNMENT.RIGHT
    )
    run = paragraph.add_run(title)
    set_font(run, size=10, bold=True)
    paragraph.add_run("\t")
    run = paragraph.add_run(dates)
    set_font(run, size=9.5)

    detail = doc.add_paragraph()
    set_paragraph_spacing(detail, after=0.4, line=1.0)
    detail.paragraph_format.keep_with_next = True
    run = detail.add_run(organization)
    set_font(run, size=9.5, italic=True)
    run = detail.add_run(f" | {location}")
    set_font(run, size=9.5)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="Resume Bullet")
    set_paragraph_spacing(paragraph, after=0.5, line=1.0)
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.17)
    fmt.first_line_indent = Inches(-0.13)
    run = paragraph.add_run("- ")
    set_font(run, size=9.7)
    run = paragraph.add_run(text)
    set_font(run, size=9.7)
    return paragraph


def add_skill_line(doc, label, text):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=0.35, line=1.0)
    run = paragraph.add_run(label)
    set_font(run, size=9.6, bold=True)
    run = paragraph.add_run(text)
    set_font(run, size=9.6)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.60)
    section.right_margin = Inches(0.60)
    section.header_distance = Inches(0.20)
    section.footer_distance = Inches(0.20)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(19)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title_style._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    if "Resume Section" not in [s.name for s in styles]:
        section_style = styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
        section_style.base_style = normal
    if "Resume Bullet" not in [s.name for s in styles]:
        bullet_style = styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet_style.base_style = normal

    doc.core_properties.title = "Samuel Geyser Northrop Grumman Internship Resume"
    doc.core_properties.author = "Samuel Geyser"
    doc.core_properties.subject = "Software engineering internship resume"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, after=1.2, line=1.0)
    run = title.add_run("Samuel Geyser")
    set_font(run, size=19, bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, after=2.8, line=1.0)
    run = contact.add_run("samgeyser@gmail.com | (970) 218-0851 | ")
    set_font(run, size=9)
    add_hyperlink(
        contact,
        "linkedin.com/in/samuel-geyser-396847393",
        "https://www.linkedin.com/in/samuel-geyser-396847393",
    )
    run = contact.add_run(" | ")
    set_font(run, size=9)
    add_hyperlink(contact, "github.com/SMGeyser", "https://github.com/SMGeyser")

    add_section_heading(doc, "Summary")
    summary = doc.add_paragraph()
    set_paragraph_spacing(summary, after=0.4, line=1.0)
    run = summary.add_run(
        "Computer Science student with a 3.99 GPA and an AI Engineering minor, graduating in December 2028. "
        "Experience evaluating approximately 50 large language model outputs per month, building applications with Go and MySQL, "
        "designing normalized MySQL databases, and leading small teams in safety-focused, time-sensitive operations."
    )
    set_font(run, size=9.7)

    add_section_heading(doc, "Education")
    education = doc.add_paragraph()
    set_paragraph_spacing(education, after=0.25, line=1.0)
    education.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.30), WD_TAB_ALIGNMENT.RIGHT
    )
    run = education.add_run("Brigham Young University-Idaho")
    set_font(run, size=10, bold=True)
    run = education.add_run(" | Rexburg, ID\tExpected Dec 2028")
    set_font(run, size=9.6)

    degree = doc.add_paragraph()
    set_paragraph_spacing(degree, after=0.25, line=1.0)
    run = degree.add_run("Bachelor of Science in Computer Science | Minor in AI Engineering | GPA 3.99/4.00")
    set_font(run, size=9.7)

    coursework = doc.add_paragraph()
    set_paragraph_spacing(coursework, after=0.35, line=1.0)
    run = coursework.add_run("Relevant Coursework  ")
    set_font(run, size=9.5, bold=True)
    run = coursework.add_run(
        "Data Science Programming, Computer Systems, Programming with Classes, AI in Software Development, "
        "Linear Algebra, SQL, Scientific Computing, Physics I and II"
    )
    set_font(run, size=9.5)

    add_section_heading(doc, "Technical Skills")
    add_skill_line(doc, "Languages  ", "Go, Python, C#, JavaScript, SQL and MySQL, HTML and CSS")
    add_skill_line(doc, "Frameworks and Libraries  ", "Flask, .NET, NumPy, Pandas, Matplotlib")
    add_skill_line(doc, "Tools and Methods  ", "Git, VS Code, Jupyter Notebook, Agile workflow management, object-oriented programming, relational database design, data analysis, LLM evaluation")

    add_section_heading(doc, "Experience")
    add_entry_heading(
        doc,
        "LLM Quality Assurance Specialist",
        "Feb 2026 - Present",
        "Handshake and DataAnnotation",
        "Remote",
    )
    add_bullet(
        doc,
        "Evaluate approximately 50 LLM outputs per month across STEM domains for factual accuracy, readability, and formatting, providing quality feedback for model improvement.",
    )
    add_bullet(
        doc,
        "Analyze visual reasoning tasks and structured data inputs to identify recurring failure patterns and gaps compared with human performance.",
    )

    add_entry_heading(
        doc,
        "Service Coordinator and Crew Chief",
        "Jul 2022 - Oct 2023",
        "Discount Tire",
        "Fort Collins, CO",
    )
    add_bullet(
        doc,
        "Led daily shop workflow and trained and supervised 3 to 6 service technicians while balancing safety, quality, and time constraints.",
    )
    add_bullet(
        doc,
        "Coordinated service for more than 20 customers per day, communicated status, and maintained customer satisfaction and safety standards.",
    )

    add_section_heading(doc, "Technical Projects")
    project = doc.add_paragraph()
    set_paragraph_spacing(project, after=0.3, line=1.0)
    run = project.add_run("Unit Log")
    set_font(run, size=10, bold=True)
    run = project.add_run(" | Go, MySQL | ")
    set_font(run, size=9.5, italic=True)
    add_hyperlink(project, "github.com/SMGeyser/UnitLog", "https://github.com/SMGeyser/UnitLog")
    add_bullet(
        doc,
        "Developing a mobile-friendly application with Go and MySQL for real-time rental property maintenance tracking, task prioritization, and team coordination.",
    )

    project = doc.add_paragraph()
    set_paragraph_spacing(project, before=0.6, after=0.3, line=1.0)
    run = project.add_run("D&D Relational Entity Repository")
    set_font(run, size=10, bold=True)
    run = project.add_run(" | MySQL | ")
    set_font(run, size=9.5, italic=True)
    add_hyperlink(
        project,
        "github.com/SMGeyser/itm220/sql_frontend",
        "https://github.com/SMGeyser/itm220/tree/main/sql_frontend",
    )
    add_bullet(
        doc,
        "Designed and deployed a relational database in third normal form to manage character attributes, equipment inventories, and campaign data.",
    )

    add_section_heading(doc, "Leadership")
    add_entry_heading(
        doc,
        "Full-time Missionary, Zone Leader and District Leader",
        "Oct 2023 - Oct 2025",
        "The Church of Jesus Christ of Latter-day Saints",
        "Sacramento, CA",
    )
    add_bullet(
        doc,
        "Served 9 months as Zone Leader and 8 months as District Leader, mentoring teams in goal setting, planning, and interpersonal problem-solving.",
    )
    add_bullet(
        doc,
        "Reviewed progress in leadership councils, implemented plans across six-week cycles, and launched weekly activities for Marshallese community members.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
